"""Synthetic workbook and CPU acceptance tests; no third-party numerical fixture."""

import hashlib
import json
from pathlib import Path
import runpy

import numpy as np
import pytest


def example(name):
    return runpy.run_path(str(Path("examples/gh4169_ultrasonic") / name))


@pytest.fixture
def workbook(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    book.create_sheet("Sheet2")
    book.create_sheet("Sheet3")
    sheet["A2"] = "Heat treatment parameters, average grain diameter and ultrasonic feature parameters of GH4169 "
    for column, value in {"A": "Sample number", "B": "Heat treatment  ", "C": "Average grain diameter/μm",
                          "D": "/(dB/mm)", "E": "/(dB/mm)", "F": "/(m/s)", "G": "/(m/s)"}.items():
        sheet[f"{column}3"] = value
    for i, identity in enumerate([f"NO.{n}" for n in range(1, 9)] + ["T1", "T2"]):
        row = i + 4
        for j, value in enumerate([identity, f"synthetic heat {i}", 10.+i*3, 0.02+i*.005, .003,
                                   6000.-i*2, .4, 111., 222., 333.], start=1):
            sheet.cell(row, j, value)
    sheet["B9"] = "11000℃/1h/WC"
    sheet["A18"] = "precomputed correlation not an input"
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    return path


def prepare(workbook, root):
    return example("prepare.py")["prepare"](workbook, root, source_spec={
        "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(), "source_kind": "input",
        "license": "Apache-2.0", "doi": "synthetic fixture", "contributors": ["test generator"],
        "url": "synthetic:workbook", "title": "synthetic schema test"})


def test_normalization_uses_only_original_columns_and_keeps_source_row(workbook, tmp_path):
    from experiment_to_cpfe.datasets.hdf5 import read_hdf5
    from experiment_to_cpfe.datasets.training import build_training_dataset
    root = tmp_path / "normalized"
    manifest = prepare(workbook, root)
    assert len(manifest["samples"]) == 10
    sample = read_hdf5(root / "NO.6.h5")
    row = sample.tables["measured_observations"][0]
    assert row["heat_treatment_raw"] == "11000℃/1h/WC"
    assert row["source_row"] == 9 and row["source_sheet"] == "Sheet1"
    assert sample.assets[0].source_kind.value == "input"
    config = json.loads((root / "build.json").read_text())
    result = build_training_dataset(config, base_dir=root)
    assert result.features.shape == (10, 2)
    assert result.features[0].tolist() == [.02, 6000.]
    assert result.targets[0] == 10.
    assert result.metadata["features"] == [{"name": "attenuation_mean", "unit": "dB/mm"},
                                           {"name": "longitudinal_velocity_mean", "unit": "m/s"}]
    assert result.groups[result.splits == "test"].tolist() == ["T1", "T2"]


@pytest.mark.parametrize("cell,value,match", [("C4", "=1+1", "formula"), ("A5", "NO.1", "identity"),
                                              ("D3", "wrong unit", "header"), ("C4", -1, "positive")])
def test_unsafe_workbook_variations_are_rejected(workbook, cell, value, match):
    import openpyxl
    book = openpyxl.load_workbook(workbook)
    book["Sheet1"][cell] = value
    book.save(workbook)
    with pytest.raises(ValueError, match=match):
        example("prepare.py")["inspect_workbook"](workbook)


def test_source_digest_mismatch_does_not_create_output(workbook, tmp_path):
    root = tmp_path / "no-output"
    with pytest.raises(ValueError, match="hash"):
        example("prepare.py")["prepare"](workbook, root)
    assert not root.exists()


def test_nested_protocol_excludes_outer_targets_from_every_inner_split():
    protocol = example("evaluate.py")["protocol"]()
    assert len(protocol["outer"]) == 4
    held = []
    for outer in protocol["outer"]:
        held.extend(outer["test"])
        assert len(outer["inner"]) == 3
        for fold in outer["inner"]:
            assert len(fold["train"]) == 4 and len(fold["validation"]) == 2
            assert not set(outer["test"]) & set(fold["train"] + fold["validation"])
            assert not set(fold["train"]) & set(fold["validation"])
    assert sorted(held) == [f"NO.{n}" for n in range(1, 9)]
    for fold in protocol["final"]:
        assert len(fold["train"]) == 6 and len(fold["validation"]) == 2
        assert fold["test"] == ["T1", "T2"]


def test_ridge_standardizes_training_only_and_ols_solves_known_relation():
    fit = example("evaluate.py")["linear_predict"]
    x = np.array([[1., 2.], [2., 0.], [3., 5.], [4., -1.]])
    y = 3+x@np.array([2., -4.])
    out = np.array([[100., -200.]])
    prediction, receipt = fit(x, y, out, alpha=0)
    np.testing.assert_allclose(prediction, 3+out@np.array([2., -4.]))
    assert receipt["x_mean"] == x.mean(axis=0).tolist()


def test_real_cpu_evaluation_and_checkpoint_readback_ignore_test_targets(workbook, tmp_path):
    pytest.importorskip("torch")
    from experiment_to_cpfe.datasets.hdf5 import read_hdf5, write_hdf5
    from experiment_to_cpfe.assets.registry import table_payload_sha256
    root = tmp_path / "normalized"
    prepare(workbook, root)
    evaluate = example("evaluate.py")["evaluate"]
    a = evaluate(root, tmp_path / "a", training_overrides={"epochs": 2, "patience": 1})
    for identity in ["T1", "T2"]:
        path = root / f"{identity}.h5"
        sample = read_hdf5(path)
        sample.tables["measured_observations"][0]["grain_diameter"] += 999.
        raw, curated = sample.assets
        digest = table_payload_sha256(sample.tables["measured_observations"])
        sample.assets = (raw, curated.model_copy(update={"sha256": digest,
            "conversion": curated.conversion.model_copy(update={"target_payload_sha256": digest})}))
        path.unlink()
        write_hdf5(sample, path)
    b = evaluate(root, tmp_path / "b", training_overrides={"epochs": 2, "patience": 1})
    assert a["readback_max_abs_error"] == 0
    assert len(a["predictions"]) == 40  # 10 independent specimens x 4 models
    for model in ["mlp", "mean", "ols", "ridge"]:
        pa = [r["prediction"] for r in a["predictions"] if r["model"] == model]
        pb = [r["prediction"] for r in b["predictions"] if r["model"] == model]
        np.testing.assert_array_equal(pa, pb)
