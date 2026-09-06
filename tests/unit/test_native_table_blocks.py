"""Synthetic native blocks catch wrong rows, evidence mixing and unit relabeling."""

from pathlib import Path

import numpy as np
import pytest

from experiment_to_cpfe.config import TabularSourceConfig, load_pipeline_config
from experiment_to_cpfe.adapters.tabular import load_tabular_source, assemble_sample


def source(path, **changes):
    data = dict(path=path, table_name="load_history", source_kind="measured",
                modality="time_series", format="txt", delimiter="\t", encoding="cp1252",
                column_map={"time": "0", "travel": "1", "force": "2", "sample_id": "3"},
                units={"time": "s", "travel": "mm", "force": "N", "sample_id": "1"},
                coordinate_frame=None, axis_order=("row",), native_layout="synthetic instrument",
                license="synthetic", block={"start_marker": "[Data]", "data_start_row": 4,
                "decimal": ",", "numeric_fields": ["time", "travel", "force"],
                "header_checks": [{"row": 1, "column": 0, "value": "Time"},
                                  {"row": 2, "column": 2, "value": "kN"}],
                "specimen": {"specimen_id": "coupon-A"}},
                conversions={"force": {"source_unit": "kN", "target_unit": "N",
                                      "factor": 1000, "offset": 0, "reason": "SI prefix conversion"}})
    data.update(changes)
    return TabularSourceConfig.model_validate(data)


@pytest.fixture
def lis(tmp_path):
    path = tmp_path / "coupon.lis"
    path.write_text("Gerät\n[Data]\nTime\tTravel\tForce\tID\ns\tmm\tkN\t1\n\n0,0\t1,2\t0,002\t001\n0,2\t1,1\t0,003\t002\n", encoding="cp1252")
    return path


def test_lis_keeps_physical_rows_ids_order_and_converts_declared_force(lis):
    rows = load_tabular_source(source(lis))
    assert [r["source_row"] for r in rows] == [6, 7]
    assert [r["sample_id"] for r in rows] == ["001", "002"]
    assert [r["force"] for r in rows] == [2, 3]
    assert [r["travel"] for r in rows] == [1.2, 1.1]


@pytest.mark.parametrize("change, match", [
    ({"header_checks": [{"row": 2, "column": 2, "value": "N"}]}, "header"),
    ({"data_start_row": 3}, "empty|numeric|missing selected cell"),
    ({"start_marker": "[Missing]"}, "marker"),
])
def test_shifted_or_wrong_block_is_rejected(lis, change, match):
    cfg = source(lis)
    payload = cfg.model_dump()
    payload["block"].update(change)
    with pytest.raises(ValueError, match=match):
        load_tabular_source(TabularSourceConfig.model_validate(payload))


def test_unit_conversion_cannot_change_target_unit_silently(lis):
    cfg = source(lis).model_dump()
    cfg["conversions"]["force"]["target_unit"] = "Pa"
    with pytest.raises(ValueError, match="unit"):
        load_tabular_source(TabularSourceConfig.model_validate(cfg))


def test_nonfinite_selected_values_rejected_without_dropping_row(lis):
    lis.write_text(lis.read_text(encoding="cp1252").replace("0,003", "NaN"), encoding="cp1252")
    with pytest.raises(ValueError, match="finite"):
        load_tabular_source(source(lis))


def test_xlsx_blocks_keep_mixed_evidence_and_separate_rows(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Curves"
    sheet.append(["Experiment", None, "Simulation", None])
    sheet.append(["strain", "stress", "strain", "stress"])
    sheet.append([0.1, 12, 0.2, 20])
    sheet.append([0.3, 14, None, None])
    path = tmp_path / "mixed.xlsx"
    book.save(path)
    base = source(path).model_dump()
    base.update(format="xlsx", delimiter=None, encoding="utf-8", conversions={},
                column_map={"strain": "0", "stress": "1"}, units={"strain": "1", "stress": "MPa"},
                block={"sheet": "Curves", "data_start_row": 3, "data_end_row": 5,
                       "numeric_fields": ["strain", "stress"]})
    exp = TabularSourceConfig.model_validate(base)
    base.update(source_kind="simulated", table_name="simulation_records",
                column_map={"strain": "2", "stress": "3"})
    base["block"]["data_end_row"] = 4
    sim = TabularSourceConfig.model_validate(base)
    cfg = load_pipeline_config(Path("examples/synthetic_minimal/sample.yaml"))
    sample = assemble_sample(cfg.model_copy(update={"sources": (exp, sim), "assets": ()}))
    assert [r["stress"] for r in sample.tables["load_history"]] == [12, 14]
    assert sample.tables["simulation_records"][0]["source_kind"] == "simulated"
    assert sample.tables["load_history"][0]["source_sheet"] == "Curves"
    assert sample.assets[1].descriptive_metadata["block"]["data_end_row"] == 5


def test_selected_xlsx_formula_is_rejected_without_using_stale_cache(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active.append(["=1+2"])
    book.save(path)
    cfg = source(path, format="xlsx", column_map={"force": "0"}, units={"force": "N"},
                 conversions={}, block={"sheet": "Sheet", "data_start_row": 1,
                                        "numeric_fields": ["force"]})
    with pytest.raises(ValueError, match="formula"):
        load_tabular_source(cfg)


def test_block_provenance_survives_hdf5_npz(lis, tmp_path):
    from experiment_to_cpfe.datasets.hdf5 import write_hdf5, read_hdf5
    from experiment_to_cpfe.datasets.package import write_npz
    import json
    cfg = load_pipeline_config(Path("examples/synthetic_minimal/sample.yaml"))
    sample = assemble_sample(cfg.model_copy(update={"sources": (source(lis),), "assets": ()}))
    out = tmp_path / "sample.h5"
    digest = write_hdf5(sample, out)
    restored = read_hdf5(out)
    child = restored.assets[1]
    assert restored.assets[0].units["force"] == "kN"
    assert child.units["force"] == "N"
    assert child.descriptive_metadata["block"]["specimen"] == {"specimen_id": "coupon-A"}
    assert child.descriptive_metadata["conversions"]["force"]["factor"] == 1000
    assert child.conversion.source_sha256 == restored.assets[0].sha256
    write_npz(out, tmp_path / "sample.npz", digest)
    with np.load(tmp_path / "sample.npz", allow_pickle=False) as payload:
        assert json.loads(payload["__tables_json__"].item()) == restored.tables
