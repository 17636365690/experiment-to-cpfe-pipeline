"""Explicit native table blocks; indices never depend on nonempty-row filtering."""

import csv
import math
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field


class HeaderCheck(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    row: int = Field(ge=1)
    column: int = Field(ge=0)
    value: str


class TableBlock(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    sheet: str | None = None
    start_marker: str | None = None
    data_start_row: int = Field(ge=1)
    data_end_row: int | None = Field(default=None, ge=1)
    decimal: Literal[".", ","] = "."
    numeric_fields: tuple[str, ...] = ()
    header_checks: tuple[HeaderCheck, ...] = ()
    specimen: dict[str, str] = Field(default_factory=dict)


class UnitConversion(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True, allow_inf_nan=False)
    source_unit: str = Field(min_length=1)
    target_unit: str = Field(min_length=1)
    factor: float
    offset: float = 0
    reason: str = Field(min_length=1)


def read_table_block(config) -> list[dict[str, object]]:
    """Read one block. Marker-relative row 1 is the first line after the marker."""
    block = config.block
    if block is None:
        raise ValueError("a block is required for native block parsing")
    if config.format not in {"csv", "txt", "xlsx"}:
        raise ValueError("blocks support CSV, TXT and XLSX only")
    if block.data_end_row is not None and block.data_end_row <= block.data_start_row:
        raise ValueError("data_end_row must be greater than data_start_row (exclusive end)")
    if not set(block.numeric_fields) <= config.column_map.keys():
        raise ValueError("numeric_fields must name mapped target columns")
    try:
        columns = {key: int(value) for key, value in config.column_map.items()}
    except ValueError as exc:
        raise ValueError("block columns must be zero-based integer indices") from exc
    if any(index < 0 for index in columns.values()):
        raise ValueError("block column indices must be nonnegative")
    if {"source_row", "source_sheet"}.intersection(columns):
        raise ValueError("source_row and source_sheet are reserved block fields")

    if config.format == "xlsx":
        if not block.sheet or block.start_marker:
            raise ValueError("XLSX requires a sheet and does not support text markers")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX requires the native extra (openpyxl)") from exc
        with config.path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                if block.sheet not in book.sheetnames:
                    raise ValueError(f"sheet is missing: {block.sheet}")
                grid = list(book[block.sheet].iter_rows(values_only=True))
            finally:
                book.close()
        offset = 0
    else:
        if block.sheet:
            raise ValueError("sheet applies only to XLSX")
        if not config.delimiter or len(config.delimiter) != 1:
            raise ValueError("native text blocks require a one-character delimiter")
        lines = config.path.read_text(encoding=config.encoding).splitlines()
        offset = 0
        if block.start_marker:
            matches = [i for i, line in enumerate(lines) if line.strip() == block.start_marker]
            if len(matches) != 1:
                raise ValueError("start marker must occur exactly once")
            offset = matches[0] + 1
        # One CSV record per physical line: multiline quoted cells need another profile.
        grid = [next(csv.reader([line], delimiter=config.delimiter, strict=True)) for line in lines]

    def cell(row, column):
        if not 0 <= row < len(grid) or not 0 <= column < len(grid[row]):
            raise ValueError(f"missing selected cell at source row {row + 1}, column {column}")
        value = grid[row][column]
        if config.format == "xlsx" and isinstance(value, str) and value.startswith("="):
            raise ValueError(f"selected formula at row {row + 1}; export evaluated values with provenance")
        return value

    for check in block.header_checks:
        if str(cell(offset + check.row - 1, check.column)).strip() != check.value:
            raise ValueError(f"header check failed at row {offset + check.row}, column {check.column}")
    stop = len(grid) if block.data_end_row is None else offset + block.data_end_row - 1
    if stop > len(grid):
        raise ValueError("selected row range exceeds the source")
    result = []
    for row_index in range(offset + block.data_start_row - 1, stop):
        record = {}
        for name, column in columns.items():
            value = cell(row_index, column)
            if value is None or (isinstance(value, str) and not value.strip()):
                raise ValueError(f"empty selected value at source row {row_index + 1}, column {column}")
            if name in block.numeric_fields:
                try:
                    value = float(str(value).replace(block.decimal, "."))
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"nonnumeric value at source row {row_index + 1}, column {column}") from exc
                if not math.isfinite(value):
                    raise ValueError(f"nonfinite value at source row {row_index + 1}, column {column}")
            else:
                value = str(value)
            record[name] = value
        record["source_row"] = row_index + 1
        if block.sheet:
            record["source_sheet"] = block.sheet
        result.append(record)
    if not result:
        raise ValueError("empty selected block")
    return result


def convert_columns(rows, config):
    """Execute declared affine conversion without dimensional or semantic inference."""
    from experiment_to_cpfe.schema.validation import _unit_is_declared
    for name, conversion in config.conversions.items():
        if (name not in config.column_map or conversion.target_unit != config.units.get(name)
                or not _unit_is_declared(conversion.source_unit)
                or not _unit_is_declared(conversion.target_unit)):
            raise ValueError(f"conversion requires mapped field and matching declared target unit: {name}")
        if conversion.factor == 0 or not conversion.reason.strip():
            raise ValueError("conversion requires a nonzero factor and a reason")
        for row in rows:
            value = row[name]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"unit conversion requires an explicitly numeric field: {name}")
            value = value * conversion.factor + conversion.offset
            if not math.isfinite(value):
                raise ValueError(f"conversion generated nonfinite value: {name}")
            row[name] = value
    return rows
