"""Normalize the pinned GH4169 workbook through the existing native table adapter."""

import argparse
import json
import math
from pathlib import Path

from experiment_to_cpfe.adapters.tabular import assemble_sample
from experiment_to_cpfe.config import PipelineConfig
from experiment_to_cpfe.datasets.hdf5 import write_hdf5
from experiment_to_cpfe.provenance.hashing import sha256_file


IDENTITIES = [f"NO.{n}" for n in range(1, 9)] + ["T1", "T2"]
UNITS = {"attenuation_mean": "dB/mm", "longitudinal_velocity_mean": "m/s", "grain_diameter": "um"}
HEADERS = {"A3": "Sample number", "B3": "Heat treatment", "C3": "Average grain diameter/μm",
           "D3": "/(dB/mm)", "F3": "/(m/s)"}


def inspect_workbook(path):
    """Read every sheet, preserve formula/cache pairs, and validate the selected block."""
    from openpyxl import load_workbook
    book = load_workbook(path, data_only=False)
    cache = load_workbook(path, data_only=True)
    try:
        sheet = book["Sheet1"]
        for address, expected in HEADERS.items():
            if str(sheet[address].value).strip() != expected:
                raise ValueError(f"unexpected header at Sheet1!{address}")
        records = []
        for row, identity in enumerate(IDENTITIES, start=4):
            if sheet[f"A{row}"].value != identity:
                raise ValueError(f"unexpected or duplicate specimen identity at row {row}")
            for column in ("C", "D", "F"):
                cell = sheet[f"{column}{row}"]
                if cell.data_type == "f":
                    raise ValueError(f"selected formula at {cell.coordinate}")
                if type(cell.value) not in (int, float) or not math.isfinite(cell.value) or cell.value <= 0:
                    raise ValueError(f"positive finite measurement required at {cell.coordinate}")
            records.append({"specimen": identity, "source_sheet": "Sheet1", "source_row": row,
                            "heat_treatment_raw": sheet[f"B{row}"].value})
        sheets = []
        for current in book:
            cells = [{"cell": cell.coordinate, "value": cell.value, "type": cell.data_type,
                      "cached": cache[current.title][cell.coordinate].value,
                      "comment": cell.comment.text if cell.comment else None}
                     for row in current for cell in row if cell.value is not None]
            sheets.append({"name": current.title, "state": current.sheet_state,
                           "dimensions": current.calculate_dimension(), "cells": cells,
                           "merged": [str(r) for r in current.merged_cells.ranges],
                           "formula_count": sum(c["type"] == "f" for c in cells)})
        return {"sheets": sheets, "records": records, "selected_numeric_columns": ["D", "F", "C"],
                "numeric_units": UNITS, "parameter_representation": "author-supplied specimen summaries",
                "excluded": {"E/G": "dispersion parameters omitted from this fixed two-mean-feature task",
                             "H/I/J": "nonlinear/backscatter definitions or units not fully resolved; header/value concerns",
                             "rows 16:49": "author correlation, mapping/fitting coefficients and evaluation results"}}
    finally:
        book.close()
        cache.close()


def prepare(source, output_dir, *, source_spec=None):
    source, output = Path(source).resolve(), Path(output_dir).resolve()
    if source_spec is None:
        source_spec = json.loads(Path(__file__).with_name("source.json").read_text(encoding="utf-8"))
    digest = sha256_file(source)
    if digest != source_spec["sha256"]:
        raise ValueError("workbook hash differs from the reviewed source")
    audit = inspect_workbook(source)
    if sha256_file(source) != digest:
        raise ValueError("workbook hash changed during inspection")
    output.mkdir(parents=True, exist_ok=False)
    inputs, samples = [], []
    for record in audit["records"]:
        identity, row = record["specimen"], record["source_row"]
        sample_config = {
            "sample": {"sample_id": identity, "experiment_id": source_spec["doi"],
                       "microstructure_id": identity, "load_path_id": "ultrasonic-characterization",
                       "schema_version": "0.1", "coordinate": {"name": "tabular_record", "axes": ["row"], "units": "1"},
                       "unit_system": UNITS, "tensor_order": ["scalar"],
                       "orientation": {"representation": "not_applicable", "reason": "scalar regression; no orientation or tensor used"}},
            "sources": [{"path": str(source), "table_name": "measured_observations", "format": "xlsx",
                         "source_kind": source_spec["source_kind"], "modality": "table", "delimiter": None,
                         "encoding": "utf-8", "column_map": {"specimen": "0", "heat_treatment_raw": "1",
                             "grain_diameter": "2", "attenuation_mean": "3", "longitudinal_velocity_mean": "5"},
                         "units": {**UNITS, "specimen": "1", "heat_treatment_raw": "1"},
                         "coordinate_frame": None, "axis_order": ["row"],
                         "native_layout": "Sheet1 specimen summaries; original columns C/D/F with identity A",
                         "license": source_spec["license"], "block": {"sheet": "Sheet1", "data_start_row": row,
                             "data_end_row": row+1, "numeric_fields": list(UNITS),
                             "header_checks": [{"row": row, "column": 0, "value": identity}],
                             "specimen": {**{k: str(v) for k, v in record.items()}, "identity_evidence": "Sheet1 column A"}}}],
            "abaqus": {"command": ["unused"], "job_name": "unused"}, "export": {"formats": ["hdf5"]}}
        sample = assemble_sample(PipelineConfig.model_validate(sample_config))
        if sample.assets[0].sha256 != digest:
            raise ValueError("workbook hash changed during normalization")
        provenance = {"source": source_spec, "selected_record": record,
                      "processing": "existing XLSX block adapter; original specimen means; no scaling or target fit"}
        write_hdf5(sample, output/f"{identity}.h5", provenance)
        (output/f"{identity}.ingest.json").write_text(json.dumps(sample_config, ensure_ascii=False, indent=2), encoding="utf-8")
        split = "test" if identity in {"T1", "T2"} else "validation" if identity in {"NO.1", "NO.5"} else "train"
        inputs.append({"path": f"{identity}.h5", "sample_id": identity, "layout": "specimen", "split": split,
                       "target_specimen": {"column": "specimen", "evidence": "Sheet1 A4:A13 sample numbers; one row per physical specimen"}})
        samples.append({**record, "path": f"{identity}.h5", "sha256": sha256_file(output/f"{identity}.h5")})
    columns = {key: {"kind": "table", "table": "measured_observations", "column": key,
                     "id_columns": ["specimen"], "source_unit": unit} for key, unit in UNITS.items()}
    config = {"version": 1, "features": [{"name": k, "unit": UNITS[k]} for k in list(UNITS)[:2]],
              "target": {"name": "grain_diameter", "unit": "um"}, "group_by": "sample_id",
              "grouping_evidence": "Ten original specimen identities; aggregate measurements kept in one group per specimen",
              "layouts": {"specimen": {"columns": columns, "alignment_evidence": "same physical worksheet row"}}, "inputs": inputs}
    manifest = {"source": {**source_spec, "local_path": str(source)}, "samples": samples}
    for name, content in [("normalization.json", manifest), ("workbook-audit.json", audit), ("build.json", config)]:
        (output/name).write_text(json.dumps(content, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    result = prepare(args.source, args.output_dir)
    print(json.dumps({"status": "completed", "specimens": len(result["samples"])}))
